#!/usr/bin/env python3
"""
Script to expand subchapters in The Midnight Confessor docx files using python-docx.
"""

from docx import Document
from docx.shared import Pt
import re
import shutil

def expand_subchapter_2_3A():
    """Expand Subchapter 2.3A - The Leverage"""
    docx_path = 'The_Midnight_Confessor_Batch1.docx'
    backup_path = 'The_Midnight_Confessor_Batch1_backup.docx'
    
    # Create backup
    shutil.copy(docx_path, backup_path)
    print(f"Created backup: {backup_path}")
    
    # Load document
    doc = Document(docx_path)
    
    # Find the narrative section for Subchapter 2.3A
    in_subchapter = False
    narrative_start_idx = None
    narrative_end_idx = None
    
    for i, para in enumerate(doc.paragraphs):
        text = para.text
        
        if 'Subchapter 2.3A' in text:
            in_subchapter = True
            continue
        
        if in_subchapter and 'NARRATIVE:' in text:
            narrative_start_idx = i
            continue
        
        if narrative_start_idx is not None:
            if '[DECISION POINT]' in text:
                narrative_end_idx = i
                break
            if 'Subchapter' in text and i > narrative_start_idx + 1:
                narrative_end_idx = i
                break
    
    if narrative_start_idx is None:
        print("ERROR: Could not find Subchapter 2.3A narrative")
        return False
    
    if narrative_end_idx is None:
        narrative_end_idx = len(doc.paragraphs)
    
    print(f"Found narrative at paragraphs {narrative_start_idx} to {narrative_end_idx}")
    
    # Expanded narrative text
    expanded_text = """NARRATIVE: "We need to find her. Now." I grabbed Silas by the arm. The time for deliberation was over.

Silas looked at the flash drive then at me. His eyes were red-rimmed and desperate. The weight of his confession hung between us like a shroud. "You have the evidence, Jack. You have me cold. You can arrest me right now. Be the hero again. It would be clean. It would be the right thing to do."

He held out his wrists. The gesture was theatrical but genuine. I looked at the cuffs on my belt. Cold steel. The tools of justice. Then I looked at the text about Maya. A girl I had never met. A girl whose mother I had put away. A girl who was now in the hands of a woman who had spent years planning this moment.

The balcony air was cold. It cut through my coat and found the old wounds. My phone buzzed again. Sarah. Another text. The FBI is mobilizing. They want Silas in custody. They want the flash drive. They want to do this the right way.

The right way. I had spent thirty years doing things the right way. Following procedure. Building cases. Trusting the system. And where had it gotten me? A partner who framed innocent men. A best friend who would turn out to be something worse. A city full of corruption I had been too blind to see.

Silas's hands were still extended. Waiting. The cuffs would click. The arrest would be clean. The conviction would be certain. And Maya Bellamy would be in the wind. Victoria would have her. The system would move too slow. The bureaucracy would grind. The paperwork would pile up. And somewhere in the city, a terrified girl would wait for a rescue that might come too late.

I looked at Silas. Really looked at him. The man I had trusted. The man who had betrayed that trust. The man who was now offering himself up as a sacrifice to the machine. "If I arrest you now, Silas, what happens to Maya?"

"She dies," he said. The words came out flat. No emotion. Just fact. "Victoria told me. She said if I was arrested before she was done with me, she would kill the girl. She said the system moves too slow. She said you would choose the law over the life. She was betting on it."

My hand moved to the cuffs. The metal was cold against my palm. The weight of thirty years of procedure. The weight of doing things the right way. The weight of being the good cop. The legend.

But legends were built on bodies. On cases closed. On convictions secured. And somewhere in those thirty years, I had stopped asking if the bodies were the right ones. I had stopped questioning if the convictions were true. I had become the machine. Efficient. Certain. Blind.

The Methodical Path demanded I secure the criminal and trust the system to find the girl. It was clean. It was legal. It was what Sarah would do. It was what the good cop would do. Secure the evidence. Build the case. Trust the process.

The Pragmatic path demanded I use the criminal to hunt the greater threat. It was dirty. It was dangerous. It was what a desperate man would do. Use the leverage. Save the girl. Ask forgiveness later.

I looked at Silas. At his extended wrists. At the flash drive on the glass table. At my phone with Sarah's text. At the city spread out below us, gray and cold and full of secrets.

The choice wasn't between right and wrong. It was between two kinds of wrong. The wrong of letting a girl die for the sake of procedure. The wrong of using a criminal to save her. Both would cost me something. Both would change who I was. But only one would let me sleep at night.

My hand moved away from the cuffs. I picked up the flash drive. I looked at Silas. "You're going to help me find her. You're going to use every contact you have. Every favor you're owed. Every dirty secret you know. And if you try to run, if you try to warn Victoria, I will make sure your family knows exactly what you did. I will make sure your sons know their father framed an innocent man. I will burn your life to the ground."

Silas's hands dropped. He nodded. "I understand."

"Do you? Because this isn't a negotiation. This is me using you. This is me becoming exactly what Victoria wants me to become. A man who breaks the rules. A man who uses criminals. A man who puts justice before the law."

"I understand," he repeated. "And I'm sorry, Jack. For everything."

The apology hung in the air between us. Empty. Meaningless. But it was all he had left to give. I turned away from him. I looked at the city. At the rain. At the gray expanse of water beyond the buildings. Somewhere out there, Maya Bellamy was waiting. And I was coming for her. With a criminal as my guide. With a flash drive as my weapon. With everything I had left to lose."""
    
    # Split into paragraphs
    expanded_paras = [p.strip() for p in expanded_text.split('\n\n') if p.strip()]
    
    # Remove old narrative paragraphs
    for i in range(narrative_end_idx - 1, narrative_start_idx - 1, -1):
        if i < len(doc.paragraphs):
            p = doc.paragraphs[i]
            p.clear()
            # Remove the paragraph element
            p_element = p._element
            p_element.getparent().remove(p_element)
    
    # Insert new paragraphs
    insert_para = doc.paragraphs[narrative_start_idx] if narrative_start_idx < len(doc.paragraphs) else None
    for para_text in expanded_paras:
        if insert_para is None:
            new_para = doc.add_paragraph(para_text)
        else:
            new_para = doc.paragraphs[narrative_start_idx].insert_paragraph_before(para_text)
        insert_para = new_para
    
    # Save
    doc.save(docx_path)
    print(f"Successfully expanded Subchapter 2.3A in {docx_path}")
    return True

if __name__ == '__main__':
    expand_subchapter_2_3A()
